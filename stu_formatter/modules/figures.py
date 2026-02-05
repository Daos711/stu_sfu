"""Figure formatting: captions, centering, spacing."""

import logging
import re

from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from ..config import STUConfig
from ..utils.xml_helpers import has_drawing, paragraph_is_empty

logger = logging.getLogger(__name__)

# Pattern for figure captions: "Рисунок 1 – Название" or "Рисунок 1.1 – Название"
FIGURE_CAPTION_RE = re.compile(
    r"^Рисунок\s+(\d+(?:\.\d+)?|[А-Я]\.\d+)\s*[–\-—]\s*(.+)",
    re.IGNORECASE,
)

# Pattern for multi-page figure captions: "Рисунок 1 – Название, лист 2"
FIGURE_SHEET_RE = re.compile(
    r"^Рисунок\s+(\d+(?:\.\d+)?|[А-Я]\.\d+)\s*(?:,\s*лист\s+\d+|[–\-—]\s*.+,\s*лист\s+\d+)",
    re.IGNORECASE,
)


def _is_figure_caption(text: str) -> bool:
    """Check if text is a figure caption."""
    stripped = text.strip()
    return bool(FIGURE_CAPTION_RE.match(stripped) or FIGURE_SHEET_RE.match(stripped))


def _is_figure_paragraph(paragraph) -> bool:
    """Check if paragraph contains only an image/drawing."""
    if not has_drawing(paragraph):
        return False
    # Check that there's no significant text besides the image
    text = paragraph.text.strip()
    return len(text) == 0


def _format_figure_caption(paragraph, config: STUConfig) -> None:
    """Format figure caption: centered, TNR 14, no indent."""
    for run in paragraph.runs:
        run.font.name = config.font_name
        run.font.size = Pt(config.font_size_figure_caption_pt)

    pf = paragraph.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.first_line_indent = Mm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def _format_image_paragraph(paragraph, config: STUConfig) -> None:
    """Format a paragraph containing an image: centered, no indent."""
    pf = paragraph.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.first_line_indent = Mm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def _ensure_blank_before(doc: Document, paragraph, index: int) -> None:
    """Ensure blank line before figure."""
    if index <= 0:
        return
    paragraphs = doc.paragraphs
    prev_p = paragraphs[index - 1]
    if not paragraph_is_empty(prev_p) and not _is_figure_paragraph(prev_p):
        from docx.oxml.ns import qn
        new_p = doc.element.makeelement(qn("w:p"), {})
        paragraph._element.addprevious(new_p)


def _ensure_blank_after(doc: Document, paragraph, index: int) -> None:
    """Ensure blank line after figure caption."""
    paragraphs = doc.paragraphs
    if index + 1 < len(paragraphs):
        next_p = paragraphs[index + 1]
        if not paragraph_is_empty(next_p):
            from docx.oxml.ns import qn
            new_p = doc.element.makeelement(qn("w:p"), {})
            paragraph._element.addnext(new_p)
    else:
        from docx.oxml.ns import qn
        new_p = doc.element.makeelement(qn("w:p"), {})
        paragraph._element.addnext(new_p)


def apply_figures(doc: Document, config: STUConfig) -> None:
    """Format figure captions and image paragraphs.

    - Image paragraphs: centered, no indent
    - Figure captions: centered, TNR 14, no indent
    - Blank lines before images and after captions
    """
    paragraphs = list(doc.paragraphs)
    figure_count = 0
    caption_count = 0

    for i, para in enumerate(paragraphs):
        text = para.text.strip()

        # Format image paragraphs
        if _is_figure_paragraph(para):
            _format_image_paragraph(para, config)
            _ensure_blank_before(doc, para, i)
            figure_count += 1
            continue

        # Format figure captions
        if _is_figure_caption(text):
            _format_figure_caption(para, config)
            _ensure_blank_after(doc, para, i)
            caption_count += 1
            continue

    logger.info(
        "Figures formatted: %d image(s), %d caption(s).",
        figure_count, caption_count,
    )

"""Basic tests for the STU formatter."""

import os
import tempfile
import unittest

from docx import Document
from docx.shared import Pt, Mm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT

from stu_formatter.config import STUConfig
from stu_formatter.formatter import STUFormatter
from stu_formatter.modules.page_setup import apply_page_setup
from stu_formatter.modules.fonts_and_paragraphs import apply_body_formatting
from stu_formatter.modules.headings import apply_headings
from stu_formatter.modules.tables import apply_tables
from stu_formatter.modules.figures import apply_figures
from stu_formatter.modules.lists import apply_lists
from stu_formatter.modules.bibliography import apply_bibliography
from stu_formatter.modules.page_numbers import apply_page_numbers
from stu_formatter.modules.formulas import apply_formulas
from stu_formatter.modules.appendices import apply_appendices

# EMU tolerance for mm-based comparisons (~1000 EMU ≈ 0.03 mm)
EMU_TOLERANCE = 1000


def _create_test_doc() -> Document:
    """Create a minimal test document with various elements."""
    doc = Document()

    # Title page
    doc.add_paragraph("Сибирский федеральный университет")
    doc.add_paragraph("Курсовая работа")
    doc.add_paragraph("")  # blank

    # Structural heading
    p_intro = doc.add_paragraph("Введение")
    p_intro.style = doc.styles["Normal"]

    doc.add_paragraph(
        "Это введение к документу. Здесь описывается актуальность работы."
    )

    # Section heading
    h1 = doc.add_heading("1 Теоретическая часть", level=1)
    doc.add_paragraph(
        "Основной текст первого раздела. Содержит описание теории."
    )

    # Subsection
    h2 = doc.add_heading("1.1 Обзор литературы", level=2)
    doc.add_paragraph(
        "Текст подраздела с обзором литературы [1, с. 45]."
    )

    # Table caption and table
    doc.add_paragraph("Таблица 1 – Результаты измерений")
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Параметр"
    table.cell(0, 1).text = "Значение"
    table.cell(0, 2).text = "Единица"
    table.cell(1, 0).text = "Длина"
    table.cell(1, 1).text = "100"
    table.cell(1, 2).text = "мм"
    table.cell(2, 0).text = "Масса"
    table.cell(2, 1).text = "50"
    table.cell(2, 2).text = "г"

    doc.add_paragraph("Текст после таблицы.")

    # Figure caption
    doc.add_paragraph("Рисунок 1 – Схема эксперимента")

    # List
    doc.add_paragraph("Перечень требований:")
    p_list1 = doc.add_paragraph("• первый пункт;")
    p_list2 = doc.add_paragraph("• второй пункт;")
    p_list3 = doc.add_paragraph("• третий пункт.")

    # Section 2
    h1_2 = doc.add_heading("2 Практическая часть", level=1)
    doc.add_paragraph("Текст второго раздела.")

    # Conclusion
    p_concl = doc.add_paragraph("Заключение")
    p_concl.style = doc.styles["Normal"]
    doc.add_paragraph("Основные выводы по работе.")

    # Bibliography
    p_bib = doc.add_paragraph("Список использованных источников")
    p_bib.style = doc.styles["Normal"]
    doc.add_paragraph("1 Иванов, И.И. Основы теории / И.И. Иванов. – М.: Наука, 2020. – 256 с.")
    doc.add_paragraph("2 Петров, П.П. Методы исследования / П.П. Петров. – СПб.: БХВ, 2021. – 312 с.")

    # Appendix
    doc.add_paragraph("ПРИЛОЖЕНИЕ А")
    doc.add_paragraph("Дополнительные данные")
    doc.add_paragraph("Здесь размещены дополнительные материалы.")

    return doc


class TestConfig(unittest.TestCase):
    """Test configuration defaults."""

    def test_default_config(self):
        config = STUConfig()
        self.assertEqual(config.font_name, "Times New Roman")
        self.assertEqual(config.font_size_pt, 14)
        self.assertEqual(config.line_spacing, 1.0)
        self.assertEqual(config.first_line_indent_mm, 12.5)
        self.assertEqual(config.margin_left_mm, 30.0)
        self.assertEqual(config.margin_right_mm, 10.0)
        self.assertEqual(config.margin_top_mm, 20.0)
        self.assertEqual(config.margin_bottom_mm, 20.0)
        self.assertFalse(config.numbering_within_sections)

    def test_custom_config(self):
        config = STUConfig(line_spacing=1.5, numbering_within_sections=True)
        self.assertEqual(config.line_spacing, 1.5)
        self.assertTrue(config.numbering_within_sections)


class TestPageSetup(unittest.TestCase):
    """Test page setup module."""

    def _assertEmu(self, actual, expected_mm, msg=None):
        """Assert EMU value is within tolerance of expected mm."""
        expected = Mm(expected_mm)
        self.assertAlmostEqual(actual, expected, delta=EMU_TOLERANCE, msg=msg)

    def test_portrait_margins(self):
        doc = Document()
        config = STUConfig()
        apply_page_setup(doc, config)

        section = doc.sections[0]
        self._assertEmu(section.page_width, 210)
        self._assertEmu(section.page_height, 297)
        self._assertEmu(section.left_margin, 30)
        self._assertEmu(section.right_margin, 10)
        self._assertEmu(section.top_margin, 20)
        self._assertEmu(section.bottom_margin, 20)


class TestBodyFormatting(unittest.TestCase):
    """Test body text formatting."""

    def test_font_and_spacing(self):
        doc = Document()
        doc.add_paragraph("Тестовый параграф с текстом.")
        config = STUConfig()
        apply_body_formatting(doc, config)

        para = doc.paragraphs[0]
        self.assertEqual(para.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        self.assertAlmostEqual(
            para.paragraph_format.first_line_indent, Mm(12.5), delta=EMU_TOLERANCE
        )
        for run in para.runs:
            self.assertEqual(run.font.name, "Times New Roman")
            self.assertEqual(run.font.size, Pt(14))

    def test_custom_spacing(self):
        doc = Document()
        doc.add_paragraph("Параграф.")
        config = STUConfig(line_spacing=1.5)
        apply_body_formatting(doc, config)

        para = doc.paragraphs[0]
        self.assertEqual(para.paragraph_format.line_spacing, 1.5)


class TestHeadings(unittest.TestCase):
    """Test heading detection and formatting."""

    def test_structural_heading(self):
        doc = Document()
        p = doc.add_paragraph("Введение")
        config = STUConfig()
        apply_headings(doc, config)

        # Should be uppercase, bold, centered
        for run in p.runs:
            self.assertTrue(run.font.bold)
            self.assertEqual(run.text, "ВВЕДЕНИЕ")
        self.assertEqual(p.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(p.paragraph_format.first_line_indent, Mm(0))

    def test_section_heading(self):
        doc = Document()
        h = doc.add_heading("1 Теоретическая часть", level=1)
        config = STUConfig()
        apply_headings(doc, config)

        for run in h.runs:
            self.assertTrue(run.font.bold)
            self.assertEqual(run.font.name, "Times New Roman")
        self.assertAlmostEqual(
            h.paragraph_format.first_line_indent, Mm(12.5), delta=EMU_TOLERANCE
        )


class TestTables(unittest.TestCase):
    """Test table formatting."""

    def test_table_caption(self):
        doc = Document()
        p = doc.add_paragraph("Таблица 1 – Результаты")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        table.cell(1, 0).text = "1"
        table.cell(1, 1).text = "2"

        config = STUConfig()
        apply_tables(doc, config)

        self.assertEqual(p.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(p.paragraph_format.first_line_indent, Mm(0))


class TestFigures(unittest.TestCase):
    """Test figure caption formatting."""

    def test_figure_caption(self):
        doc = Document()
        p = doc.add_paragraph("Рисунок 1 – Схема")
        config = STUConfig()
        apply_figures(doc, config)

        self.assertEqual(p.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(p.paragraph_format.first_line_indent, Mm(0))


class TestLists(unittest.TestCase):
    """Test list formatting."""

    def test_bullet_replacement(self):
        doc = Document()
        p = doc.add_paragraph("• первый пункт")
        config = STUConfig()
        apply_lists(doc, config)

        # Bullet should be replaced with dash
        self.assertTrue(p.runs[0].text.startswith("–"))


class TestBibliography(unittest.TestCase):
    """Test bibliography formatting."""

    def test_bibliography_section(self):
        doc = Document()
        doc.add_paragraph("Список использованных источников")
        entry = doc.add_paragraph("1 Иванов, И.И. Книга. – М.: Наука, 2020.")
        config = STUConfig()
        apply_bibliography(doc, config)

        self.assertEqual(entry.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        self.assertAlmostEqual(
            entry.paragraph_format.first_line_indent, Mm(12.5), delta=EMU_TOLERANCE
        )


class TestFullFormat(unittest.TestCase):
    """Integration test: format a complete test document."""

    def test_full_document(self):
        doc = _create_test_doc()

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp_in:
            doc.save(tmp_in.name)
            input_path = tmp_in.name

        output_path = input_path.replace(".docx", "_formatted.docx")

        try:
            config = STUConfig()
            formatter = STUFormatter(config)
            formatter.format(input_path, output_path)

            # Verify output exists
            self.assertTrue(os.path.exists(output_path))

            # Open and verify basic properties
            result = Document(output_path)
            section = result.sections[0]
            self.assertAlmostEqual(section.page_width, Mm(210), delta=EMU_TOLERANCE)
            self.assertAlmostEqual(section.page_height, Mm(297), delta=EMU_TOLERANCE)
            self.assertAlmostEqual(section.left_margin, Mm(30), delta=EMU_TOLERANCE)
            self.assertAlmostEqual(section.right_margin, Mm(10), delta=EMU_TOLERANCE)
        finally:
            os.unlink(input_path)
            if os.path.exists(output_path):
                os.unlink(output_path)


class TestCLIArgs(unittest.TestCase):
    """Test CLI argument parsing."""

    def test_default_args(self):
        from stu_formatter.main import parse_args
        args = parse_args(["test.docx"])
        self.assertEqual(args.input, "test.docx")
        self.assertIsNone(args.output)
        self.assertEqual(args.spacing, 1.0)
        self.assertFalse(args.section_numbering)
        self.assertFalse(args.verbose)

    def test_custom_args(self):
        from stu_formatter.main import parse_args
        args = parse_args([
            "input.docx",
            "-o", "output.docx",
            "--spacing", "1.5",
            "--section-numbering",
            "-v",
        ])
        self.assertEqual(args.input, "input.docx")
        self.assertEqual(args.output, "output.docx")
        self.assertEqual(args.spacing, 1.5)
        self.assertTrue(args.section_numbering)
        self.assertTrue(args.verbose)


if __name__ == "__main__":
    unittest.main()
